"""Refresh data/denial_cheatsheet.json from the billing SOP workbook.

Only two sheets are read — "Denial Cheatsheet" and "Appeals". The workbook
also holds the team's portal logins ("Insurance Portal"); nothing from any
other sheet is copied, and the workbook itself is never committed.

    python3 scripts/import_denial_sop.py ~/Downloads/SOP.xlsx
"""
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(os.path.dirname(HERE), 'data', 'denial_cheatsheet.json')
SHEETS = ('Denial Cheatsheet', 'Appeals')


def _text(v):
    return '' if v is None else str(v).strip()


def read_sop(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    names = {n.lower(): n for n in wb.sheetnames}
    out = {'source': os.path.basename(path), 'rules': [], 'appeals': []}
    sheet = wb[names.get('denial cheatsheet') or names['denial cheat sheet']]
    rows = [r for r in sheet.iter_rows(values_only=True) if any(_text(c) for c in r)]
    for r in rows[1:]:   # S No · Codes · Drug Name · Payer · Reason · Action
        r = list(r) + [None] * 6
        if not _text(r[1]) and not _text(r[4]):
            continue
        out['rules'].append({'sop_row': _text(r[0]), 'codes': _text(r[1]), 'drug': _text(r[2]),
                             'payer': _text(r[3]), 'reason': _text(r[4]), 'action': _text(r[5])})
    if 'appeals' in names:
        last = ''
        for r in wb[names['appeals']].iter_rows(values_only=True):
            r = list(r) + [None, None]
            title, text = _text(r[0]), _text(r[1])
            if not text or text.lower().startswith('appeal / dispute'):
                continue
            last = title or last
            out['appeals'].append({'title': last, 'text': text})
    return out


def main(argv=None):
    argv = argv if argv is not None else sys.argv[1:]
    if not argv:
        print(__doc__)
        return 2
    data = read_sop(argv[0])
    with open(OUT, 'w', encoding='utf-8') as fh:
        json.dump(data, fh, indent=1, ensure_ascii=False)
    print(f"{len(data['rules'])} rule(s), {len(data['appeals'])} appeal text(s) → {OUT}")
    return 0


if __name__ == '__main__':
    sys.exit(main())
